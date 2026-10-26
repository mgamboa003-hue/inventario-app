"""Movimientos: altas, ajustes, bajas, saldo, destino, filtros y Excel."""
import io
import sqlite3

import db


def _filas(sql, params=()):
    c = sqlite3.connect(db.SQLITE_PATH)
    c.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in c.execute(sql, params).fetchall()]
    finally:
        c.close()


def _crear(cliente, nombre, codigo, stock):
    cliente.post("/productos/nuevo", data={"nombre": nombre, "codigo": codigo,
                                           "stock_minimo": "1", "stock_actual": str(stock)})
    return _filas("SELECT id FROM productos WHERE codigo = ?", (codigo,))[0]["id"]


def _editar(cliente, pid, stock, motivo=""):
    return cliente.post(f"/productos/{pid}/editar", data={
        "nombre": "Rep Ajuste", "codigo": "AJ-1", "stock_minimo": "1",
        "stock_actual": str(stock), "planta": "quilicura", "motivo_ajuste": motivo,
    }, follow_redirects=True)


def test_crear_repuesto_registra_alta_con_saldo(admin_client):
    pid = _crear(admin_client, "Rep Alta", "ALTA-1", 7)
    movs = _filas("SELECT tipo, cantidad, saldo FROM movimientos WHERE producto_id = ?", (pid,))
    assert movs == [{"tipo": "alta", "cantidad": 7, "saldo": 7}]
    html = admin_client.get("/movimientos").get_data(as_text=True)
    assert "Rep Alta" in html and "Alta" in html


def test_ajuste_de_stock_exige_motivo_y_queda_como_movimiento(admin_client):
    pid = _crear(admin_client, "Rep Ajuste", "AJ-1", 5)
    r = _editar(admin_client, pid, 8)
    assert "motivo del ajuste" in r.get_data(as_text=True).lower()
    assert _filas("SELECT stock_actual FROM productos WHERE id = ?", (pid,))[0]["stock_actual"] == 5

    _editar(admin_client, pid, 8, "conteo físico")
    assert _filas("SELECT stock_actual FROM productos WHERE id = ?", (pid,))[0]["stock_actual"] == 8
    aj = _filas("SELECT cantidad, saldo, motivo FROM movimientos WHERE producto_id = ? AND tipo = 'ajuste'", (pid,))
    assert aj == [{"cantidad": 3, "saldo": 8, "motivo": "conteo físico"}]


def test_editar_sin_cambiar_stock_no_pide_motivo_ni_crea_ajuste(admin_client):
    pid = _crear(admin_client, "Rep Ajuste", "AJ-1", 5)
    _editar(admin_client, pid, 5)
    assert not _filas("SELECT 1 FROM movimientos WHERE producto_id = ? AND tipo = 'ajuste'", (pid,))


def test_salida_guarda_saldo_y_destino(admin_client):
    pid = _crear(admin_client, "Rep Salida", "SAL-1", 10)
    admin_client.post("/movimientos/nuevo/salida", data={
        "producto_id": pid, "cantidad": "4", "usuario": "Tester", "destino": "Selladora 2",
    }, follow_redirects=True)
    m = _filas("SELECT cantidad, saldo, destino FROM movimientos WHERE producto_id = ? AND tipo = 'salida'", (pid,))
    assert m == [{"cantidad": 4, "saldo": 6, "destino": "Selladora 2"}]
    assert "Selladora 2" in admin_client.get("/movimientos").get_data(as_text=True)


def test_baja_y_restaurado_quedan_en_movimientos(admin_client):
    pid = _crear(admin_client, "Rep Baja", "BAJA-1", 2)
    admin_client.post(f"/productos/{pid}/eliminar")
    admin_client.post(f"/productos/{pid}/restaurar")
    tipos = [r["tipo"] for r in _filas("SELECT tipo FROM movimientos WHERE producto_id = ? ORDER BY id", (pid,))]
    assert tipos == ["alta", "baja", "restaurado"]


def test_filtro_por_tipo_y_usuario(admin_client):
    pid = _crear(admin_client, "Rep Filtro", "FIL-1", 10)
    admin_client.post("/movimientos/nuevo/salida", data={"producto_id": pid, "cantidad": "1", "usuario": "Tester Filtro"})
    html = admin_client.get("/movimientos?tipo=alta").get_data(as_text=True)
    assert "1 movimiento" in html
    html = admin_client.get("/movimientos?usuario=tester filtro ").get_data(as_text=True)
    assert "1 movimiento" in html and "Tester Filtro" in html


def test_exportar_movimientos_a_excel(admin_client):
    from openpyxl import load_workbook
    pid = _crear(admin_client, "Rep Excel", "XLS-1", 3)
    admin_client.post("/movimientos/nuevo/salida", data={"producto_id": pid, "cantidad": "1", "destino": "Prensa"})
    r = admin_client.get("/movimientos/exportar")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.data)).active
    filas = list(ws.iter_rows(values_only=True))
    assert filas[0][:6] == ("Fecha", "Tipo", "Código", "Repuesto", "Cantidad", "Saldo")
    assert ("Salida", -1, 2, "Prensa") in [(f[1], f[4], f[5], f[7]) for f in filas[1:]]


def test_migracion_crea_saldo_inicial_y_rellena_saldos(admin_client):
    """Datos 'viejos': stock sin alta y movimientos sin saldo -> la migracion los cuadra."""
    c = sqlite3.connect(db.SQLITE_PATH)
    c.execute("INSERT INTO productos (nombre, codigo, stock_actual, planta) VALUES ('Viejo', 'OLD-1', 9, 'quilicura')")
    pid = c.execute("SELECT id FROM productos WHERE codigo = 'OLD-1'").fetchone()[0]
    c.execute("INSERT INTO movimientos (producto_id, tipo, cantidad, fecha) VALUES (?, 'entrada', 5, '2026-09-01 10:00:00')", (pid,))
    c.execute("INSERT INTO movimientos (producto_id, tipo, cantidad, fecha) VALUES (?, 'salida', 2, '2026-09-02 10:00:00')", (pid,))
    c.execute("DELETE FROM sistema_meta WHERE clave = 'migracion_saldos_movimientos_v1'")
    c.commit(); c.close()

    conn = db.get_db_connection()
    db._migrar_saldos_movimientos(conn)
    db._migrar_saldos_movimientos(conn)   # segunda vez no hace nada
    conn.close()

    movs = _filas("SELECT tipo, cantidad, saldo FROM movimientos WHERE producto_id = ? ORDER BY fecha, id", (pid,))
    assert movs == [
        {"tipo": "saldo_inicial", "cantidad": 6, "saldo": 6},
        {"tipo": "entrada", "cantidad": 5, "saldo": 11},
        {"tipo": "salida", "cantidad": 2, "saldo": 9},
    ]
